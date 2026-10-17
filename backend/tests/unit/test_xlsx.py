"""Mục O: chuỗi ghi vào Excel luôn là ô chữ, không thành công thức, không thêm dấu '."""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from app.services.xlsx import append_row, to_bytes


def test_strings_stay_plain_text() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    values = [
        '=HYPERLINK("https://evil.example","Mở")',
        "+84",
        "-Ab3_x9-Tm",
        "@x",
        "Shop Thỏ",
        5,
    ]
    append_row(ws, values)
    sheet = load_workbook(io.BytesIO(to_bytes(wb))).active
    assert sheet is not None
    cells = next(sheet.iter_rows(min_row=1, max_row=1))
    assert [c.value for c in cells] == values
    assert cells[0].data_type == "s"
