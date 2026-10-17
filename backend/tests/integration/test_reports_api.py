"""Tổng quan, Fanpage, Chi tiết Trang, Quảng cáo (Bước 7): số liệu, phạm vi vai trò, xuất Excel."""

from __future__ import annotations

import datetime as dt
import io
from decimal import ROUND_HALF_UP, Decimal
from typing import Any
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import AdInsightDaily, Campaign, Lead
from app.services.report_export import safe_cell
from tests.integration.report_fixtures import FISS, RABBIT, SMART, SOFTWARE, Demo

pytestmark = pytest.mark.integration


@pytest.mark.parametrize(
    ("role", "path", "code"),
    [
        ("s1", "/api/reports/overview", 403),
        ("s1", "/api/reports/pages", 403),
        ("s1", "/api/reports/ads", 403),
        ("lead", "/api/reports/ads", 403),
        ("lead", "/api/reports/overview", 200),
        ("viewer", "/api/reports/ads", 200),
        ("viewer", "/api/reports/pages", 200),
        ("manager", "/api/reports/ads", 200),
    ],
)
async def test_report_permissions(as_user: Any, role: str, path: str, code: int) -> None:
    resp = await (await as_user(role)).get(path)
    assert resp.status_code == code, resp.text


async def test_overview_numbers_for_owner(as_user: Any, demo: Demo) -> None:
    body = (await (await as_user("owner")).get("/api/reports/overview?range=7d")).json()
    k = body["kpis"]
    assert Decimal(k["new_leads"]["value"]) == 22
    assert Decimal(k["close_rate"]["value"]) == Decimal("0.2727")  # 6 / 22
    assert Decimal(k["revenue_vnd"]["value"]) == 10_000_000
    assert Decimal(k["phone_rate"]["value"]) == Decimal("0.5909")  # 13 / 22
    funnel = body["funnel"]
    assert funnel["new"] == 22 and funnel["won"] == 6 and funnel["phone"] == 13
    assert sum(d["new"] for d in body["daily"]) == 22 and len(body["daily"]) == 7

    # Chi phí / khách chốt = chi tiêu QC 7 ngày ÷ khách chốt từ QC (4).
    campaigns = {c["id"]: c for c in body["campaigns"]}
    fiss = campaigns[demo.campaign_id]
    assert fiss["leads"] == 12 and fiss["won"] == 4
    assert fiss["cost_per_won_vnd"] == round(fiss["spend_vnd"] / 4)
    spend = sum(c["spend_vnd"] for c in body["campaigns"] if c["spend_vnd"] is not None)
    # Tài khoản USD chưa có tỉ giá: chỉ cộng phần chi tiêu đã quy được VND.
    assert Decimal(k["cost_per_won_vnd"]["value"]) == round(Decimal(spend) / 4)
    # Tài khoản USD chưa có tỉ giá.
    assert body["missing_fx"] == ["USD"]

    alerts = body["alerts"]
    assert alerts["unassigned_pages"] == 1
    assert alerts["bm_disconnected"] == 0
    assert body["is_mock"] is True
    # Mục M: "Dẫn đầu" chỉ hạng A/B (cao trước), "Cần hỗ trợ" chỉ hạng C/D (thấp trước).
    top, bottom = body["boards"]["top"], body["boards"]["bottom"]
    assert all(b["tier"] in ("A", "B") for b in top)
    assert all(b["tier"] in ("C", "D") for b in bottom)
    top_totals = [Decimal(b["total"]) for b in top]
    bottom_totals = [Decimal(b["total"]) for b in bottom]
    assert top_totals == sorted(top_totals, reverse=True)
    assert bottom_totals == sorted(bottom_totals)
    assert {b["member"]["name"] for b in top + bottom} == {
        demo.accounts["s1"].user.full_name,
        demo.accounts["s2"].user.full_name,
    }


async def test_overview_team_lead_sees_only_team_without_ad_costs(as_user: Any) -> None:
    body = (await (await as_user("lead")).get("/api/reports/overview?range=7d")).json()
    assert Decimal(body["kpis"]["new_leads"]["value"]) == 10  # chỉ khách của s1
    assert body["kpis"]["cost_per_won_vnd"] is None
    assert body["campaigns"] is None
    assert body["alerts"]["unassigned_pages"] is None
    assert body["alerts"]["expensive_campaigns"] is None
    assert [t["name"] for t in body["teams"]] == ["[Nhóm 1]"]


async def test_overview_team_filter_and_forbidden_team(as_user: Any) -> None:
    owner = await as_user("owner")
    teams = {t["name"]: t["id"] for t in (await owner.get("/api/reports/overview")).json()["teams"]}
    body = (await owner.get(f"/api/reports/overview?team_id={teams['[Nhóm 2]']}")).json()
    assert Decimal(body["kpis"]["new_leads"]["value"]) == 12
    lead = await as_user("lead")
    resp = await lead.get(f"/api/reports/overview?team_id={teams['[Nhóm 2]']}")
    assert resp.status_code == 403


async def test_custom_range_is_limited(as_user: Any) -> None:
    owner = await as_user("owner")
    ok = await owner.get("/api/reports/overview?range=custom&from=2026-09-01&to=2026-09-30")
    assert ok.status_code == 200 and len(ok.json()["daily"]) == 30
    too_long = await owner.get("/api/reports/overview?range=custom&from=2026-01-01&to=2026-09-30")
    assert too_long.status_code == 422
    assert too_long.json()["detail"]["code"] == "range_too_long"


async def test_pages_tabs_and_scope(as_user: Any) -> None:
    body = (await (await as_user("manager")).get("/api/reports/pages")).json()
    assert body["total"] == 4 and body["tabs"]["all"] == 4
    assert body["tabs"]["unassigned"] == 1
    assert body["tabs"]["declining"] == 1  # Mua Sắm Thông Minh giảm lượt xem
    rows = {r["fb_page_id"]: r for r in body["rows"]}
    assert rows[FISS]["leads_7d"] == 12 and rows[FISS]["won_7d"] == 4
    assert rows[SMART]["owner"] is None and rows[SMART]["status"] == "declining"
    assert rows[SOFTWARE]["views_30d"] > 0 and len(rows[SOFTWARE]["views_daily"]) == 30
    assert body["rows"][0]["fb_page_id"] == FISS  # mặc định: nhiều khách nhất
    assert body["can_assign"] is True

    unassigned = (await (await as_user("manager")).get("/api/reports/pages?tab=unassigned")).json()
    assert [r["fb_page_id"] for r in unassigned["rows"]] == [SMART]
    found = (await (await as_user("manager")).get("/api/reports/pages?q=fiss")).json()
    assert [r["fb_page_id"] for r in found["rows"]] == [FISS]

    team = (await (await as_user("lead")).get("/api/reports/pages")).json()
    assert {r["fb_page_id"] for r in team["rows"]} == {SOFTWARE, RABBIT}
    assert team["can_assign"] is False and team["bms"] == []


async def test_page_detail(as_user: Any, demo: Demo) -> None:
    manager = await as_user("manager")
    body = (await manager.get(f"/api/reports/pages/{demo.pages[SOFTWARE].id}")).json()
    assert body["owner"]["name"] == demo.accounts["s1"].user.full_name
    assert body["team"] == "[Nhóm 1]"
    assert len(body["daily_views"]) == 30
    assert Decimal(body["kpis"]["views"]["value"]) > 0
    assert body["kpis"]["followers"] > 0
    posts = body["posts"]
    assert posts and posts[0]["created_time"] >= posts[-1]["created_time"]
    # Phần Mềm đăng khoảng 6 bài/tuần: đủ bài để vẽ bảng nhiệt 7 × 6.
    assert body["heatmap"] is not None and len(body["heatmap"]) == 7
    assert all(len(row) == 6 for row in body["heatmap"])

    lead = await as_user("lead")
    assert (await lead.get(f"/api/reports/pages/{demo.pages[FISS].id}")).status_code == 404
    assert (await lead.get(f"/api/reports/pages/{demo.pages[RABBIT].id}")).status_code == 200


async def test_ads_report_levels_and_account_filter(as_user: Any, demo: Demo) -> None:
    viewer = await as_user("viewer")
    body = (await viewer.get("/api/reports/ads?range=7d")).json()
    rows = {r["id"]: r for r in body["rows"]}
    fiss = rows[demo.campaign_id]
    assert fiss["leads"] == 12 and fiss["won"] == 4
    assert fiss["results"] > 0 and fiss["cost_per_result_vnd"] > 0
    assert fiss["daily_budget_vnd"] == 450_000
    assert body["kpis"]["won"]["value"] == "4"
    assert body["missing_fx"] == ["USD"]
    assert body["cost_threshold_vnd"] == 25_000
    assert len(body["accounts"]) == 2

    vnd = next(a for a in body["accounts"] if a["currency"] == "VND")
    only_vnd = (await viewer.get(f"/api/reports/ads?range=7d&account_id={vnd['id']}")).json()
    assert only_vnd["missing_fx"] == ["USD"]  # cảnh báo chung, không theo tài khoản
    assert all(r["account"] == vnd["name"] for r in only_vnd["rows"])

    ads = (await viewer.get("/api/reports/ads?range=7d&level=ad")).json()
    assert all(r["parent"] for r in ads["rows"])
    adsets = (await viewer.get("/api/reports/ads?range=7d&level=adset")).json()
    assert sum(r["leads"] for r in adsets["rows"]) == 12


async def test_paused_campaign_without_spend_is_hidden(
    as_user: Any, admin_db: AsyncSession
) -> None:
    paused = await admin_db.scalar(
        select(Campaign.id).where(Campaign.fb_campaign_id == "23800000000000004")
    )
    body = (await (await as_user("owner")).get("/api/reports/ads?range=7d")).json()
    assert str(paused) not in {r["id"] for r in body["rows"]}


async def test_exports(as_user: Any) -> None:
    owner = await as_user("owner")
    resp = await owner.get("/api/reports/pages.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml"
    )
    sheet = load_workbook(io.BytesIO(resp.content)).active
    assert sheet is not None
    names = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert "FISS" in names and len(names) == 4

    overview = await owner.get("/api/reports/overview.xlsx?range=30d")
    book = load_workbook(io.BytesIO(overview.content))
    assert book.sheetnames == ["Tong quan", "Chien dich"]

    lead = await as_user("lead")
    lead_book = load_workbook(io.BytesIO((await lead.get("/api/reports/overview.xlsx")).content))
    assert lead_book.sheetnames == ["Tong quan"]  # không có số tiền QC


def test_safe_cell_blocks_formulas() -> None:
    assert safe_cell("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert safe_cell("+84") == "'+84"
    assert safe_cell("Shop Thỏ") == "Shop Thỏ"
    assert safe_cell(Decimal("1.5")) == 1.5


async def test_lead_detail_shows_estimated_cost_only_with_ads_permission(
    as_user: Any, demo: Demo, admin_db: AsyncSession
) -> None:
    lead = await admin_db.scalar(
        select(Lead)
        .where(Lead.page_id == demo.pages[FISS].id)
        .order_by(Lead.first_inbound_at.desc())
        .limit(1)
    )
    assert lead is not None
    vn = ZoneInfo("Asia/Ho_Chi_Minh")
    day = lead.first_inbound_at.astimezone(vn).date()
    spend = await admin_db.scalar(
        select(AdInsightDaily.spend_vnd).where(
            AdInsightDaily.level == "campaign",
            AdInsightDaily.object_id == lead.campaign_id,
            AdInsightDaily.date == day,
        )
    )
    start = dt.datetime.combine(day, dt.time(0), tzinfo=vn)
    same_day = await admin_db.scalar(
        select(func.count()).where(
            Lead.campaign_id == lead.campaign_id,
            Lead.first_inbound_at >= start,
            Lead.first_inbound_at < start + dt.timedelta(days=1),
        )
    )
    assert spend and same_day

    owner = (await (await as_user("owner")).get(f"/api/leads/{lead.id}")).json()
    assert owner["campaign_name"] == "[Tin nhắn - Giải pháp tài chính]"
    assert owner["cost_visible"] is True
    expected = (Decimal(spend) / same_day).quantize(Decimal(1), rounding=ROUND_HALF_UP)
    assert owner["estimated_cost_vnd"] == int(expected)

    staff = (await (await as_user("s2")).get(f"/api/leads/{lead.id}")).json()
    assert staff["campaign_name"] == "[Tin nhắn - Giải pháp tài chính]"
    assert staff["cost_visible"] is False and staff["estimated_cost_vnd"] is None
