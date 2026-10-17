"""Nhập nhân viên và gán Trang bằng Excel (Bước 6, decisions.md mục H).

Cột: Họ tên · Email · Nhóm · Vai trò (Nhân viên/Trưởng nhóm) · ID hoặc tên Trang.
- Xem trước: đọc file, kiểm tra từng dòng, lưu `import_jobs` (không lưu file gốc).
- Nhập: kiểm tra lại với dữ liệu hiện tại, chỉ nhập dòng hợp lệ. Tài khoản mới nhận mật khẩu tạm
  ngẫu nhiên (bắt đổi lần đầu); file mật khẩu tạm trả về một lần, không lưu.
"""

from __future__ import annotations

import asyncio
import base64
import datetime as dt
import io
import re
import secrets
import unicodedata
import uuid
from typing import TYPE_CHECKING, Any

from email_validator import EmailNotValidError, validate_email
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import hash_password
from app.models import (
    ImportJob,
    MemberRole,
    Page,
    PageAssignment,
    Team,
    User,
    WorkspaceMember,
    WorkspaceSettings,
)
from app.schemas.imports import ImportPreview, ImportResult, ImportRow
from app.services import audit
from app.services.assignment import assign_page, end_assignment
from app.services.auth_service import normalize_email
from app.services.errors import ServiceError

if TYPE_CHECKING:
    from app.api.deps import AuthContext

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
HEADERS = ("Họ tên", "Email", "Nhóm", "Vai trò", "ID hoặc tên Trang")
ROLE_NAMES = {
    "": MemberRole.STAFF,
    "nhan vien": MemberRole.STAFF,
    "truong nhom": MemberRole.TEAM_LEAD,
}
HASH_CONCURRENCY = 4


def _fold(text: str) -> str:
    """So khớp không phân biệt hoa thường, dấu và khoảng trắng thừa."""
    text = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", text).strip().lower()


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # ID Trang nhập dạng số
    return str(value).strip()


# ---------------------------------------------------------------- đọc file


def read_rows(data: bytes) -> list[tuple[int, list[str]]]:
    if not data:
        raise ServiceError("import_file_empty", 422)
    if len(data) > MAX_BYTES:
        raise ServiceError("import_file_too_large", 413)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ServiceError("import_file_invalid", 422) from None
    sheet = workbook.worksheets[0]
    rows: list[tuple[int, list[str]]] = []
    header: list[str] | None = None
    for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        cells = [_cell(v) for v in values[:5]] + [""] * (5 - len(values[:5]))
        if header is None:
            header = [_fold(c) for c in cells]
            if header[:2] != [_fold(HEADERS[0]), _fold(HEADERS[1])]:
                raise ServiceError("import_header_invalid", 422)
            continue
        if not any(cells):
            continue
        rows.append((index, cells))
        if len(rows) > MAX_ROWS:
            raise ServiceError("import_too_many_rows", 422, max=MAX_ROWS)
    workbook.close()
    if header is None:
        raise ServiceError("import_header_invalid", 422)
    return rows


# ---------------------------------------------------------------- kiểm tra


def _email(raw: str) -> str:
    """Email chuẩn hóa, hoặc chuỗi rỗng nếu không hợp lệ."""
    try:
        return normalize_email(validate_email(raw, check_deliverability=False).normalized)
    except EmailNotValidError:
        return ""


async def _validate(db: AsyncSession, raw: list[tuple[int, list[str]]]) -> list[ImportRow]:
    pages = list(await db.scalars(select(Page).where(Page.is_tracked.is_(True))))
    by_id = {p.fb_page_id: p for p in pages}
    by_name: dict[str, list[Page]] = {}
    for p in pages:
        by_name.setdefault(_fold(p.name), []).append(p)
    owners = {
        page_id: email
        for page_id, email in (
            await db.execute(
                select(PageAssignment.page_id, User.email)
                .join(WorkspaceMember, WorkspaceMember.id == PageAssignment.member_id)
                .join(User, User.id == WorkspaceMember.user_id)
                .where(PageAssignment.valid_to.is_(None))
            )
        ).all()
    }
    members = {
        email: role
        for email, role in (
            await db.execute(
                select(User.email, WorkspaceMember.role).join(
                    WorkspaceMember, WorkspaceMember.user_id == User.id
                )
            )
        ).all()
    }
    # Mục O: email đã có tài khoản nhưng không thuộc workspace này (RLS chỉ cho thấy membership
    # của workspace này) thì không nhập.
    file_emails = {e for _, cells in raw if (e := _email(cells[1]))}
    elsewhere = set(await db.scalars(select(User.email).where(User.email.in_(file_emails)))) - set(
        members
    )
    seen_emails: set[str] = set()
    seen_pages: set[uuid.UUID] = set()
    team_leads: set[str] = set()
    out: list[ImportRow] = []
    for row_no, (name, email_raw, team, role_raw, page_raw) in raw:
        errors: list[str] = []
        if not name:
            errors.append("name_required")
        elif len(name) > 200:
            errors.append("name_too_long")
        email = _email(email_raw)
        if not email:
            errors.append("email_invalid")
        elif email in elsewhere:
            errors.append("email_unavailable")
        if email and email in seen_emails:
            errors.append("email_duplicate_in_file")
        seen_emails.add(email)
        role = ROLE_NAMES.get(_fold(role_raw))
        if role is None:
            errors.append("role_invalid")
        existing_role = members.get(email)
        if existing_role in (MemberRole.OWNER, MemberRole.MANAGER, MemberRole.VIEWER):
            errors.append("cannot_edit_role")
        if role == MemberRole.TEAM_LEAD:
            if not team:
                errors.append("team_required_for_lead")
            elif _fold(team) in team_leads:
                errors.append("team_lead_duplicate")
            else:
                team_leads.add(_fold(team))
        page: Page | None = None
        if page_raw:
            matches = [by_id[page_raw]] if page_raw in by_id else by_name.get(_fold(page_raw), [])
            if not matches:
                errors.append("page_not_found")
            elif len(matches) > 1:
                errors.append("page_name_ambiguous")
            else:
                page = matches[0]
                current = owners.get(page.id)
                if current is not None and current != email:
                    errors.append("page_has_owner")
                if page.id in seen_pages:
                    errors.append("page_duplicate_in_file")
                seen_pages.add(page.id)
        out.append(
            ImportRow(
                row=row_no,
                full_name=name,
                email=email or email_raw,
                team=team or None,
                role=role,
                page=page.name if page else (page_raw or None),
                page_id=page.id if page else None,
                action=None if errors else ("update" if email in members else "create"),
                errors=errors,
            )
        )
    return out


async def preview(db: AsyncSession, ctx: AuthContext, data: bytes, file_name: str) -> ImportPreview:
    rows = await _validate(db, read_rows(data))
    ok = sum(1 for r in rows if not r.errors)
    job = ImportJob(
        workspace_id=ctx.workspace_id,
        type="members",
        file_name=file_name[:200] or "nhan-vien.xlsx",
        status="previewed",
        rows_total=len(rows),
        rows_ok=ok,
        rows_failed=len(rows) - ok,
        report={"rows": [r.model_dump(mode="json") for r in rows]},
        created_by_member_id=ctx.member_id,
    )
    db.add(job)
    await db.commit()
    return ImportPreview(
        id=job.id,
        file_name=job.file_name,
        rows_total=len(rows),
        rows_ok=ok,
        rows_failed=len(rows) - ok,
        rows=rows,
    )


# ---------------------------------------------------------------- nhập


def _temp_password() -> str:
    # 12 ký tự ngẫu nhiên + hậu tố: luôn qua chính sách mật khẩu (≥ 10 ký tự, không phổ biến).
    return secrets.token_urlsafe(9) + "-Tm"


def _credentials_xlsx(rows: list[tuple[str, str, str]]) -> str:
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 — Workbook mới luôn có một sheet
    ws.title = "Mat khau tam"
    ws.append(["Họ tên", "Email", "Mật khẩu tạm (bắt đổi ở lần đăng nhập đầu)"])
    for row in rows:
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode()


async def apply(
    db: AsyncSession, ctx: AuthContext, job_id: uuid.UUID, *, now: dt.datetime | None = None
) -> ImportResult:
    now = now or dt.datetime.now(dt.UTC)
    job = await db.scalar(select(ImportJob).where(ImportJob.id == job_id).with_for_update())
    if job is None or job.type != "members":
        raise ServiceError("import_not_found", 404)
    if job.status != "previewed":
        raise ServiceError("import_already_applied", 409)
    stored = [ImportRow.model_validate(r) for r in job.report["rows"]]
    # Kiểm tra lại: dữ liệu có thể đã đổi từ lúc xem trước.
    rows = await _validate(
        db,
        [
            (r.row, [r.full_name, r.email, r.team or "", _role_text(r.role), r.page or ""])
            for r in stored
        ],
    )
    valid = [r for r in rows if not r.errors]
    new_rows = [r for r in valid if r.action == "create"]
    passwords = {r.email: _temp_password() for r in new_rows}
    sem = asyncio.Semaphore(HASH_CONCURRENCY)

    async def _hash(pw: str) -> str:
        async with sem:
            return await asyncio.to_thread(hash_password, pw)

    hashes = dict(
        zip(passwords, await asyncio.gather(*(_hash(pw) for pw in passwords.values())), strict=True)
    )
    teams = {_fold(t.name): t for t in await db.scalars(select(Team))}
    settings = await db.get(WorkspaceSettings, ctx.workspace_id)
    actor = await db.get(WorkspaceMember, ctx.member_id)
    credentials: list[tuple[str, str, str]] = []
    created = updated = 0
    for r in valid:
        team = None
        if r.team:
            team = teams.get(_fold(r.team))
            if team is None:
                team = Team(workspace_id=ctx.workspace_id, name=r.team.strip()[:120])
                db.add(team)
                await db.flush()
                teams[_fold(r.team)] = team
        user = await db.scalar(select(User).where(User.email == r.email))
        if user is None:
            user = User(
                email=r.email,
                full_name=r.full_name,
                password_hash=hashes[r.email],
                must_change_password=True,
            )
            db.add(user)
            await db.flush()
            credentials.append((r.full_name, r.email, passwords[r.email]))
        member = await db.scalar(select(WorkspaceMember).where(WorkspaceMember.user_id == user.id))
        role = r.role or MemberRole.STAFF
        if member is None and r.email not in passwords:
            continue  # Mục O: tài khoản của workspace khác, `_validate` đã chặn; phòng khi lọt.
        if member is None:
            member = WorkspaceMember(workspace_id=ctx.workspace_id, user_id=user.id, role=role)
            db.add(member)
            created += 1
        else:
            member.role = role
            updated += 1
        member.team_id = team.id if team else None
        await db.flush()
        if team is not None and role == MemberRole.TEAM_LEAD:
            team.lead_member_id = member.id
        if r.page_id is not None:
            page = await db.get(Page, r.page_id)
            assert page is not None  # noqa: S101 — vừa kiểm tra lại ở trên
            if settings and settings.one_page_per_member:
                # Chuyển Trang: kết thúc Trang cũ của người này trước (nhập lại file để cập nhật).
                for old in await db.scalars(
                    select(Page)
                    .join(PageAssignment, PageAssignment.page_id == Page.id)
                    .where(
                        PageAssignment.member_id == member.id,
                        PageAssignment.valid_to.is_(None),
                        Page.id != page.id,
                    )
                ):
                    await end_assignment(db, old, now)
            await assign_page(db, page, member, now, assigned_by=actor, reason="[Nhập Excel]")
        await audit.record(
            db,
            workspace_id=ctx.workspace_id,
            actor_member_id=ctx.member_id,
            action="member.imported",
            entity_type="member",
            entity_id=member.id,
            after={
                "email": r.email,
                "role": role.value,
                "team": r.team,
                "page_id": str(r.page_id) if r.page_id else None,
            },
            ip=ctx.ip,
        )
    job.status = "applied"
    job.applied_at = now
    job.rows_ok = len(valid)
    job.rows_failed = len(rows) - len(valid)
    job.report = {"rows": [r.model_dump(mode="json") for r in rows]}
    await db.commit()
    return ImportResult(
        created=created,
        updated=updated,
        skipped=len(rows) - len(valid),
        credentials_file=_credentials_xlsx(credentials) if credentials else None,
    )


def _role_text(role: MemberRole | None) -> str:
    if role == MemberRole.TEAM_LEAD:
        return "Trưởng nhóm"
    return "Nhân viên" if role == MemberRole.STAFF else "?"


async def template(db: AsyncSession) -> bytes:
    """File mẫu: sheet nhập liệu và sheet danh sách Trang đang theo dõi (để chép ID)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101
    ws.title = "Nhan vien"
    ws.append(list(HEADERS))
    ws.append(["[Nhân viên mẫu 1]", "nhanvien1@example.com", "[Nhóm mẫu]", "Trưởng nhóm", ""])
    ws.append(["[Nhân viên mẫu 2]", "nhanvien2@example.com", "[Nhóm mẫu]", "Nhân viên", ""])
    for column, width in zip("ABCDE", (28, 32, 20, 16, 30), strict=True):
        ws.column_dimensions[column].width = width
    pages = wb.create_sheet("Trang")
    pages.append(["ID Trang", "Tên Trang"])
    for page in await db.scalars(select(Page).where(Page.is_tracked.is_(True)).order_by(Page.name)):
        pages.append([page.fb_page_id, page.name])
    pages.column_dimensions["A"].width = 22
    pages.column_dimensions["B"].width = 40
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
