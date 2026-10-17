"""Xuất Excel cho Tổng quan và Fanpage (Bước 7). Không có SĐT hay nội dung tin nhắn."""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from openpyxl import Workbook

from app.i18n import vi
from app.schemas.reports import OverviewOut, PageReportRow

_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    """Tên Trang, chiến dịch đến từ Facebook: chặn Excel hiểu thành công thức."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIX):
        return "'" + value
    if isinstance(value, Decimal):
        return float(value)
    return value


def _save(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _percent(value: Decimal | None) -> float | None:
    return round(float(value) * 100, 1) if value is not None else None


def overview_xlsx(data: OverviewOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 — Workbook mới luôn có một sheet
    ws.title = vi.EXPORT_OVERVIEW_SHEET
    ws.append([vi.EXPORT_RANGE.format(first=data.range.first_day, last=data.range.last_day)])
    ws.append(vi.EXPORT_OVERVIEW_HEADER)
    for key, label in vi.EXPORT_OVERVIEW_ROWS.items():
        metric = getattr(data.kpis, key)
        if metric is None:
            continue
        convert = _percent if key in ("phone_rate", "close_rate") else safe_cell
        ws.append([label, convert(metric.value), convert(metric.previous)])
    if data.campaigns is not None:
        sheet = wb.create_sheet(vi.EXPORT_CAMPAIGN_SHEET)
        sheet.append(vi.EXPORT_CAMPAIGN_HEADER)
        for c in data.campaigns:
            sheet.append([safe_cell(c.name), c.spend_vnd, c.leads, c.won, c.cost_per_won_vnd])
    return _save(wb)


def pages_xlsx(rows: list[PageReportRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 — Workbook mới luôn có một sheet
    ws.title = vi.EXPORT_PAGES_SHEET
    ws.append(vi.EXPORT_PAGES_HEADER)
    for r in rows:
        ws.append(
            [
                safe_cell(r.name),
                r.fb_page_id,
                safe_cell(r.bm_name),
                safe_cell(r.owner.name) if r.owner else None,
                safe_cell(r.owner_total),
                r.owner_tier.value if r.owner_tier else None,
                r.leads_7d,
                round(r.median_response_seconds / 60) if r.median_response_seconds else None,
                r.won_7d,
                r.posts_7d,
                r.views_30d,
                vi.EXPORT_PAGE_STATUS[r.status],
            ]
        )
    return _save(wb)
